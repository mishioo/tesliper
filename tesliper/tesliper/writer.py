###################
###   IMPORTS   ###
###################

import csv
import numpy as np
import logging as lgg
import os
import openpyxl as oxl
from collections import defaultdict

##################
###   LOGGER   ###
##################

logger = lgg.getLogger(__name__)
logger.setLevel(lgg.DEBUG)

    
###################
###   CLASSES   ###
###################

class Writer:

    def __init__(self, tesliper):
        self.ts = tesliper
        self.path = None
    
    @property
    def distribution_center(self):
        return dict(
            energies = {'txt': self.energies_txt,
                        'csv': self.energies_csv,
                        'xlsx': self.energies_xlsx},
            bars = {'txt': self.bars_txt,
                    'csv': self.bars_csv,
                    'xlsx': self.bars_xlsx},
            spectra = {'txt': self.spectra_txt,
                       'csv': self.spectra_csv,
                       'xlsx': self.spectra_xlsx},
            averaged = {'txt': self.averaged_txt,
                        'csv': self.averaged_csv,
                        'xlsx': self.averaged_xlsx}
            )
    
    def save_output(self, output, format=None, output_dir=None):
        #populations, bars (with e-m), spectra, averaged, settings
        # if 'popul' in args:
            # for en in self.energies.values():
                # path = os.path.join(output_dir,
                                    # 'Distribution.{}.txt'.format(en.type))
                # f = open(path, 'w', newline='')
                # writer = csv.writer(f, delimiter='\t')
                # writer.writerow(['Gaussian output file', 'Population', 'DE',
                                 # 'Energy', 'Imag', 'Stoichiometry'])
                # writer.writerows([[f, p, d, e, i, s] for f, p, d, e, i, s in \
                    # zip(en.filenames, en.populations, en.deltas, en.values,
                        # self.bars.iri.imag.sum(0), en.stoich)])
                # f.close()
                
        output_dir = output_dir if output_dir else self.ts.output_dir
        self.path = output_dir
        format = ['txt'] if not format else [format] \
            if not isinstance(format, (list, tuple)) else format
        output = output if isinstance(output, (list, tuple)) else [output]
        functions = []
        for thing in output:
            for fmt in format:
                try:
                    functions.append(self.distribution_center[thing][fmt])
                except KeyError:
                    logger.error('Can not export {} as {}. No such '
                        'thing or unsupported format.'.format(thing, fmt))
        for func in functions: func()
        
        # if 'ens' in args:
            # if 'txt' in format:
                # self.ens_txt_separately()
                # self.ens_txt_collectively()
            # if 'xlsx' in format:
                # self.energies_xlsx()
            # if 'csv' in format:
                # self.energies_csv()
        # if 'bars' in args:
            # self.bars_txts()
        # if 'spectra' in args:
            # for spc in self.spectra.values():
                # spc.export_txts()
            # logger.info("Individual conformers' spectra text export done.")
                    
        # if 'averaged' in args:
            # pass
        # if 'settings' in args:
            # pass
    
    __header = dict(
        rot = 'Rot. Str. ',
        dip = 'Dip. Str. ',
        roa1 = 'ROA1      ',
        raman1 = 'Raman1    ',
        vrot = 'Rot.(velo)',
        lrot = 'Rot. (len)',
        vosc = 'Osc.(velo)',
        losc = 'Osc. (len)',
        iri = 'IR Int.   ',
        vemang = 'E-M Angle ',
        eemang = 'E-M Angle ',
        zpe = 'Zero-point',
        ten = 'Thermal',
        ent = 'Enthalpy',
        gib = 'Gibbs',
        scf = 'SCF',
        ex_en = 'Excit. Energy',
        freq = 'Frequency ',
        wave = 'Wavelenght'
        )
    
    __formatters = dict(
        rot = '{:> 10.4f}',
        dip = '{:> 10.4f}',
        roa1 = '{:> 10.4f}',
        raman1 = '{:> 10.4f}',
        vrot = '{:> 10.4f}',
        lrot = '{:> 10.4f}',
        vosc = '{:> 10.4f}',
        losc = '{:> 10.4f}',
        iri = '{:> 10.4f}',
        vemang = '{:> 10.4f}',
        eemang = '{:> 10.4f}',
        zpe = '{:> 13.4f}',
        ten = '{:> 13.4f}',
        ent = '{:> 13.4f}',
        gib = '{:> 13.4f}',
        scf = '{:> 13.4f}',
        ex_en = '{:> 13.4f}',
        freq = '{:> 10.2f}',
        wave = '{:> 10.2f}'
        )
        
    energies_order = 'zpe ten ent gib scf'.split(' ')
    
    def energies_txt(self):
        self.ens_txt_collectively()
        self.ens_txt_separately()
    
    def ens_txt_separately(self):
        h = ' | '.join(['Population / %', 'Min. B. Factor',
                       'DE / (kcal/mol)', 'Energy / Hartree', 'Imag'])
        for key, en in self.ts.energies.items():
            max_fnm = max(np.vectorize(len)(en.filenames).max(), 20)
            max_stoich = max(np.vectorize(len)(en.stoich).max(), 13)
            file_path = os.path.join(self.path, '!distribution.{}.txt'.format(key))
            header = '{:<{w}} | '.format('Gaussian output file', w=max_fnm) + h
            header = header + ' | {:<{w}}'.format('Stoichiometry', w=max_stoich)
            with open(file_path, 'w') as file:
                file.write(header + '\n')
                file.write('-' * len(header) + '\n')
                for row in zip(en.filenames, en.populations * 100,
                               en.min_factor, en.deltas, en.values,
                               self.ts.bars.iri.imag.sum(0), en.stoich):
                    row = ['{:{a}{w}{f}}'.format(v, a=a, w=w, f=f) \
                        for v, a, w, f in zip(row, 
                            ('<', '^', '^', '^', '^', '^', '^'),
                            (max_fnm, 14, 14, 15, 16, 4, max_stoich),
                            ('', '.4f', '.4f', '.4f', 'f', 'd', ''))]
                    file.write(' | '.join(row) + '\n')
        logger.info('Energies separate export to text files done.')
        
    def ens_txt_collectively(self):
        with self.ts.unified_data(data_type='e') as data:
            ens = [data[en] for en in self.energies_order]
                #get them sorted
            filenames = ens[0].filenames
            longest = max(np.vectorize(len)(filenames).max(), 20)
            types = [en.type for en in ens]
            values = np.array([en.values for en in ens]).T
            #deltas = np.array([en.deltas for en in ens])
            popul = np.array([en.populations * 100 for en in ens]).T
            header = '{:<{lgst}} | {:^50} | {:^70}'.format(
                'Gaussian output file', 'Population / %', 'Energy / Hartree',
                lgst=longest)
            names = [self.__header[en] for en in self.energies_order]
        with open(os.path.join(self.path, '!distribution.txt'), 'w') as file:
            file.write(header + '\n')
            names_line = ' ' * longest + ' | ' + \
                '  '.join(['{:<{w}}'.format(n, w=max(8, len(n))) \
                           for n in names]) + ' | ' + \
                '  '.join(['{:<{w}}'.format(n, w=14 if n=='SCF' else 12) \
                           for n in names]) + '\n'
            file.write(names_line)
            file.write('-' * len(header) + '\n')
            for fnm, vals, pops in zip(filenames, values, popul):
                p_line = '  '.join(
                    ['{:>{w}.4f}'.format(p, w=max(8, len(n))) \
                     for p, n in zip(pops, names)])
                v_line = '  '.join(
                    ['{:> {w}.{prec}f}'.format(v, w=14 if n=='SCF' else 12,
                                               prec=8 if n=='SCF' else 6) \
                     for v, n in zip(vals, names)])
                line = '{:<{w}}'.format(fnm, w=longest) + ' | ' + p_line + \
                    ' | ' + v_line + '\n'
                file.write(line)
        logger.info('Energies collective export to text file done.')

    def energies_xlsx(self):
        wb = oxl.Workbook()
        ws = wb.active
        ws.title = 'Collective overview'
        ws['A1'] = 'Gaussian output file'
        ws['B1'] = 'Populations / %'
        ws['G1'] = 'Energies / hartree'
        names = [self.__header[name] for name in self.energies_order]
        ws.append([''] + names + names)
        ws.merge_cells('A1:A2')
        ws.merge_cells('B1:F1')
        ws.merge_cells('G1:K1')
        ws.freeze_panes = 'A3'
        with self.ts.unified_data(data_type='e') as data:
            filenames = data['scf'].filenames
            values = [data[name].values for name in self.energies_order]
            populs = [data[name].populations for name in self.energies_order]
            for row in zip(filenames, *populs, *values):
                ws.append(row)
        iri = self.ts.bars.iri
        for name in self.energies_order:
            ws = wb.create_sheet(title=self.__header[name])
            en = self.ts.energies[name]
            iri.trimmer.match(en)
            ws.freeze_panes = 'A2'
            header = ['Gaussian output file', 'Population / %',
                      'Min. B. Factor', 'DE / (kcal/mol)', 'Energy / Hartree',
                      'Imag', 'Stoichiometry']
            ws.append(header)
            for row in zip(en.filenames, en.populations, en.min_factor,
                           en.deltas, en.values, iri.imag.sum(0), en.stoich):
                ws.append(row)
        wb.save(os.path.join(self.path, '!distribution.xlsx'))
        logger.info('Energies export to xlsx files done.')
        
    def energies_csv(self):
        header = 'population min_factor delta energy'.split(' ')
        header = ['Gaussian output file'] + header
        for name, en in self.ts.energies.items():
            file_path = os.path.join(self.path,
                '!distribution.{}.csv'.format(name))
            with open(file_path, 'w', newline='') as file:
                csvwriter = csv.writer(file)
                csvwriter.writerow(header)
                for row in zip(en.filenames, en.populations, en.min_factor,
                               en.deltas, en.values):
                    csvwriter.writerow(row)
        logger.info('Energies export to csv files done.')
        
    def bars_export(self):
        separated = defaultdict(list)
        for bar in self.ts.bars.values():
            separated[bar._soxhlet_id].append(bar)
        order = 'dip rot raman1 roa1 vemang vrot vosc lrot losc eemang'\
                .split(' ')
        sox_ref = {'=vcd': 'vibra',
                   '=roa': 'raman',
                   'td=': 'electr'}
        for sox_id, bars in separated.items():
            com = self.ts.soxhlet.instances[sox_id].command
            _type = [val for key, val in sox_ref.items() if key in com][0]
            bars_sorted = \
                [bar for name in order for bar in bars if bar.type == name]
            freq_type = 'wave' if _type == 'electr' else 'freq'
            values_sorted = [iter(bar.values) for bar in bars_sorted]
            frequencies = iter(bars[0].full.frequencies)
            logger.debug('Will make an attempt to export data to txt '
                                 'from soxhlet {}.'.format(sox_id))
            if not bars_sorted:
                logger.debug('This soxhlet instance have not provided any '
                             'exportable data. Continuing to next soxhlet.')
                continue
            logger.debug('This soxhlet instance provided following data'
                         ' types: {}.'.format(', '.join([bar.type for bar \
                                                         in bars_sorted])))
            for fname in self.ts.soxhlet.instances[sox_id].gaussian_files:
                values, types = zip(
                    *[(next(val), bar.type) for val, bar
                      in zip(values_sorted, bars_sorted)
                      if fname in bar.filenames])
                freqs = [next(frequencies)]
                if values:
                    values = freqs + list(values)
                    types = [freq_type] + list(types)
                    yield ('.'.join(fname.split('.')[:-1]), _type, types,
                           np.array(values).T)
                    
    def bars_txt(self):
        for fname, _type, types, values_list in self.bars_export():
            filename = '{}.{}.bar.txt'.format(fname, _type)
            with open(os.path.join(self.path, filename), 'w') as file:
                file.write('\t'.join([self.__header[type] for type in types]))
                file.write('\n')
                for values in values_list:
                    line = '\t'.join(self.__formatters[tp].format(v) \
                        for v, tp in zip(values, types))
                    file.write(line + '\n')
        logger.info('Bars export to text files done.')

    def bars_csv(self):
        for fname, _type, types, values_list in self.bars_export():
            file_path = (os.path.join(self.path, 
                '{}.{}.bar.csv'.format(fname, _type)))
            with open(file_path, 'w', newline='') as file:
                csvwriter = csv.writer(file)
                csvwriter.writerow([self.__header[name] for name in types])
                for values in values_list:
                    csvwriter.writerow(values)
        logger.info('Bars export to csv files done.')
                
    def bars_xlsx(self):
        wbs = defaultdict(oxl.Workbook)
        for fname, _type, types, values_list in self.bars_export():
            wb = wbs[fname]
            ws = wb.create_sheet()
            ws.title = _type
            ws.freeze_panes = 'A2'
            ws.append([self.__header[name] for name in types])
            for values in values_list:
                ws.append([*values])
        for fname, wb in wbs.items():
            del wb['Sheet']
            wb.save(os.path.join(self.path, fname + '.bar.xlsx'))
        logger.info('Bars export to xlsx files done.')
        
    @property
    def exported_filenames(self):
        make_new_names = lambda fnm: \
            '{}.{}.txt'.format('.'.join(fnm.split('.')[:-1]), self.name)
        names = map(make_new_names, self.filenames)
        return names
    
    @property
    def averaged_filename(self):
        return 'avg_{}_{}'.format(self.name, self.energy_type)
    
    def spectra_export(self, format=''):
        for spectra in self.ts.spectra.values():
            if format == 'xlsx':
                yield spectra
            else:
                for fnm, values in zip(spectra.filenames, spectra.values):
                    filename = '.'.join(fnm.split('.')[:-1])
                    yield (filename, spectra.name, spectra, values, spectra.base,
                           spectra.text_header)
                
    def spectra_txt(self):
        for fnm, name, spc, values, base, title in self.spectra_export():
            file_path = os.path.join(self.path, '{}.{}.txt'.format(fnm, name))
            with open(file_path, 'w') as file:
                file.write(title + '\n')
                file.write(
                    '\n'.join(
                    '{:>4d}\t{: .4f}'.format(int(b), s) \
                    for b, s in zip(base, values))
                    )
        logger.info('Spectra export to text files done.')

    def spectra_csv(self): 
        for fnm, name, spc, values, base, title in self.spectra_export():
            file_path = os.path.join(self.path, '{}.{}.csv'.format(fnm, name))
            with open(file_path, 'w', newline='') as file:
                csvwriter = csv.writer(file)
                for row in zip(base, values): csvwriter.writerow(row)
        logger.info('Spectra export to csv files done.')
    
    def spectra_xlsx(self):
        wb = oxl.Workbook()
        del wb['Sheet']
        for spectra in self.spectra_export('xlsx'):
            ws = wb.create_sheet()
            ws.title = spectra.name
            ws.freeze_panes = 'B2'
            A0 = 'Wavelenght' if spectra.name in ('uv','ecd') else 'Frequency'
            ws.append([A0] + list(spectra.filenames))
            ws["A1"].comment = oxl.comments.Comment(spectra.text_header,
                                                    'Tesliper')
            for line in zip(spectra.base, *spectra.values):
                ws.append(line)
        wb.save(os.path.join(self.path, '!spectra.xlsx'))
        logger.info('Spectra export to xlsx file done.')

    def averaged_export(self):
        for name, spc in self.ts.spectra.items():
            spectra = [spc.average(self.ts.energies[en]) \
                       for en in self.energies_order]
            freq_type = 'wave' if spc.type == 'electr' else 'freq'
            yield (name, spc, spectra, freq_type)
        
    def averaged_txt(self):
        for name, spc, averaged, freq_type in self.averaged_export():
            for en, av in zip(self.energies_order, averaged):
                file_path = os.path.join(self.path, 
                    'averaged.{}.{}.txt'.format(name, en))
                with open(file_path, 'w') as file:
                    file.write(spc.text_header + '\n')
                    file.write('{} conformes averaged based on {}\n'.format(
                        len(spc._averaged[en]['populations']),
                        self.__header[en])
                        )
                    file.write(
                        '\n'.join(
                        '{:>4d}\t{: .4f}'.format(int(b), s) \
                        for b, s in zip(*av))
                        )
        logger.info('Averaged export to text files done.')
        
    def averaged_csv(self):
        for name, spc, averaged, freq_type in self.averaged_export():
            for en, av in zip(self.energies_order, averaged):
                file_path = os.path.join(self.path, 
                    'averaged.{}.{}.csv'.format(name, en))
                with open(file_path, 'w', newline='') as file:
                    csvwriter = csv.writer(file)
                    for row in zip(*av): csvwriter.writerow(row)
        logger.info('Averaged export to csv files done.')
    
    def averaged_xlsx(self):
        for name, spc, averaged, freq_type in self.averaged_export():
            wb = oxl.Workbook()
            del wb['Sheet']
            for en, av in zip(self.energies_order, averaged):
                ws = wb.create_sheet()
                ws.title = self.__header[en]
                for row in zip(*av): ws.append(row)
            wb.save(os.path.join(self.path,'averaged.{}.xlsx'.format(name,en)))
        logger.info('Averaged export to xlsx files done.')
